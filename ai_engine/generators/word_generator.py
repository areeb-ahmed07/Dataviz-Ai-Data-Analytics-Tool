"""
Word Report Generator - creates professional styled Word (.docx) documents.

Consumes ExecutiveReportData enriched with real analysis results.
Sections: Executive Summary, KPIs, Quality, Statistics, Correlations,
Outliers, ML Models, Trends, Anomalies, Recommendations.

Single Responsibility: converts ExecutiveReportData into a formatted Word document.
"""

from __future__ import annotations

import os
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from ai_engine.domain.entities import ExecutiveReportData


def _cell_bg(cell, color_hex: str):
    try:
        cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'))
    except Exception:
        pass


def _cell_pad(cell, top=80, bottom=80, left=120, right=120):
    try:
        cell._tc.get_or_add_tcPr().append(
            parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        )
    except Exception:
        pass


class WordGenerator:
    """Generates a comprehensive Word document from ExecutiveReportData."""

    def __init__(self, data: ExecutiveReportData) -> None:
        self._data = data

    def generate(self, output_path: str) -> Optional[str]:
        try:
            doc = Document()
            for s in doc.sections:
                s.top_margin = Inches(1.0)
                s.bottom_margin = Inches(1.0)
                s.left_margin = Inches(1.0)
                s.right_margin = Inches(1.0)

            DARK = RGBColor(15, 23, 42)
            MUTED = RGBColor(100, 116, 139)
            TEXT = RGBColor(30, 41, 59)
            BLUE = RGBColor(79, 70, 229)

            d = self._data

            # Title
            p = doc.add_paragraph()
            r = p.add_run(d.title)
            r.font.size = Pt(22)
            r.font.bold = True
            r.font.color.rgb = DARK

            # Metadata
            meta = doc.add_paragraph()
            mr = meta.add_run(
                f"Generated: {d.generated_at.strftime('%Y-%m-%d %H:%M:%S')}  |  "
                f"Dataset: {d.dataset_name} ({d.dataset_shape[0]:,} rows x {d.dataset_shape[1]} cols)  |  "
                f"Quality: {d.quality_score:.1f}/100"
            )
            mr.font.size = Pt(9)
            mr.font.color.rgb = MUTED

            doc.add_paragraph("\u2500" * 68).paragraph_format.space_after = Pt(8)

            # Helper for section headings
            def heading(text):
                h = doc.add_paragraph()
                h.paragraph_format.space_before = Pt(14)
                h.paragraph_format.space_after = Pt(6)
                hr = h.add_run(text)
                hr.font.size = Pt(13)
                hr.font.bold = True
                hr.font.color.rgb = BLUE

            def body(text):
                bp = doc.add_paragraph()
                bp.paragraph_format.space_after = Pt(6)
                bp.add_run(text)

            def styled_table(headers, rows, col_widths=None):
                tbl = doc.add_table(rows=1, cols=len(headers))
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for i, h in enumerate(headers):
                    c = tbl.rows[0].cells[i]
                    c.text = h
                    _cell_bg(c, "0F172A")
                    _cell_pad(c)
                    for run in c.paragraphs[0].runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9)
                for r_idx, row in enumerate(rows):
                    row_cells = tbl.add_row().cells
                    for c_idx, val in enumerate(row):
                        row_cells[c_idx].text = str(val)
                        _cell_pad(row_cells[c_idx])
                        if r_idx % 2 == 1:
                            _cell_bg(row_cells[c_idx], "F8FAFC")
                        for run in row_cells[c_idx].paragraphs[0].runs:
                            run.font.size = Pt(8.5)
                return tbl

            # 1. EXECUTIVE SUMMARY
            heading("1. Executive Summary")
            body(d.summary_paragraph or f"Comprehensive analysis of {d.dataset_name}.")

            # KPIs
            if d.kpis:
                heading("2. Key Performance Indicators")
                styled_table(
                    ["Metric", "Value", "Status"],
                    [[k.get("title", ""), k.get("value", "0"), k.get("status", "OK")] for k in d.kpis]
                )
                doc.add_paragraph().paragraph_format.space_after = Pt(6)

            # 3. DATA QUALITY
            qd = d.quality_detail
            if qd:
                heading("3. Data Quality Assessment")
                body(f"Quality Score: {qd.get('quality_score', d.quality_score):.1f}/100 | Missing: {d.missing_percentage:.2f}% | Duplicates: {qd.get('duplicate_rows', 0):,}")
                missing = qd.get("missing_per_column", [])
                if missing:
                    styled_table(
                        ["Column", "Missing Count", "Missing %"],
                        [[m["column"], str(m["missing_count"]), f"{m['missing_pct']:.2f}%"] for m in missing[:15]]
                    )
                doc.add_paragraph().paragraph_format.space_after = Pt(6)

            # 4. STATISTICS
            stats = d.statistics
            ns = stats.get("numeric_stats", [])
            if ns:
                heading("4. Descriptive Statistics")
                styled_table(
                    ["Column", "Mean", "Std Dev", "Min", "Median", "Max", "Skewness"],
                    [[s["column"], f"{s['mean']:.4f}", f"{s['std']:.4f}", f"{s['min']:.4f}", f"{s['median']:.4f}", f"{s['max']:.4f}", f"{s['skewness']:.4f}"] for s in ns[:20]]
                )
                doc.add_paragraph().paragraph_format.space_after = Pt(6)

            cs = stats.get("categorical_stats", [])
            if cs:
                heading("5. Categorical Analysis")
                cat_rows = []
                for c in cs[:15]:
                    top = c.get("top_values", [{}])[0] if c.get("top_values") else {}
                    cat_rows.append([c["column"], str(c.get("unique", 0)), str(top.get("value", "N/A")), f"{top.get('frequency', 0):.1f}%"])
                styled_table(["Column", "Unique", "Top Value", "Frequency"], cat_rows)
                doc.add_paragraph().paragraph_format.space_after = Pt(6)

            # 6. CORRELATIONS
            pairs = d.correlations.get("top_pairs", []) if d.correlations else []
            if pairs:
                heading("6. Correlation Analysis")
                styled_table(
                    ["Feature A", "Feature B", "Correlation", "Strength"],
                    [[p["col_a"], p["col_b"], f"{p.get('correlation', 0):.4f}", "Strong" if abs(p.get('correlation', 0)) > 0.8 else "Moderate"] for p in pairs[:15]]
                )
                doc.add_paragraph().paragraph_format.space_after = Pt(6)

            # 7. OUTLIERS
            iqr = d.outliers.get("iqr_outliers", []) if d.outliers else []
            if iqr:
                heading("7. Outlier Detection")
                styled_table(
                    ["Column", "Count", "%", "Lower Bound", "Upper Bound"],
                    [[o["column"], str(o.get("count", 0)), f"{o.get('percentage', 0):.1f}%", f"{o.get('lower_bound', 0):.4f}", f"{o.get('upper_bound', 0):.4f}"] for o in iqr[:15]]
                )
                doc.add_paragraph().paragraph_format.space_after = Pt(6)

            # 8. TRENDS
            if d.trends:
                heading("8. AI Trend Analysis")
                styled_table(
                    ["Column", "Direction", "R2", "Change %", "Summary"],
                    [[t.column, t.direction.value.upper(), f"{t.r2_score:.3f}", f"{t.pct_change:+.1f}%", t.summary_text[:60]] for t in d.trends[:10]]
                )
                doc.add_paragraph().paragraph_format.space_after = Pt(6)

            # 9. ANOMALIES
            if d.anomalies and d.anomalies.anomalous_rows_count > 0:
                heading("9. Anomaly Assessment")
                body(f"{d.anomalies.anomalous_rows_count:,} anomalous rows ({d.anomalies.anomaly_percentage:.1f}%).")
                for item in d.anomalies.detected_anomalies[:8]:
                    bp = doc.add_paragraph(style='List Bullet')
                    bp.add_run(f"Row #{item.row_index} | '{item.column}' = {item.value}: {item.reason}")

            # 10. ML MODELS
            if d.ml_models:
                heading("10. ML Model Performance")
                for m in d.ml_models:
                    mp = doc.add_paragraph()
                    mr = mp.add_run(f"{m['model_name']} ({m['algorithm']}) - {m.get('problem_type', '').title()}")
                    mr.font.bold = True
                    mr.font.size = Pt(11)
                    metrics = m.get("metrics", {})
                    if metrics:
                        metric_str = " | ".join([f"{k}: {v:.4f}" if isinstance(v, float) else f"{k}: {v}" for k, v in metrics.items()])
                        mi = doc.add_paragraph()
                        mir = mi.add_run(metric_str)
                        mir.font.size = Pt(9.5)
                        mir.font.color.rgb = MUTED

            # 11. RECOMMENDATIONS
            if d.recommendations:
                heading("11. Strategic Recommendations")
                for idx, rec in enumerate(d.recommendations, 1):
                    rp = doc.add_paragraph()
                    rr = rp.add_run(f"{idx}. [{rec.impact.value.upper()}] {rec.title} ({rec.category})")
                    rr.font.bold = True
                    rr.font.size = Pt(11)

                    dp = doc.add_paragraph()
                    dr = dp.add_run(rec.description)
                    dr.font.italic = True
                    dr.font.size = Pt(10)

                    if rec.action_items:
                        for a in rec.action_items:
                            ap = doc.add_paragraph(style='List Bullet 2')
                            ap.add_run(a).font.size = Pt(9.5)

            os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
            doc.save(output_path)
            return output_path
        except Exception as e:
            print(f"[WordGenerator] Error: {e}")
            import traceback
            traceback.print_exc()
            return None
