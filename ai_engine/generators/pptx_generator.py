"""
PowerPoint Report Generator - creates executive PowerPoint (.pptx) presentations.

Consumes ExecutiveReportData enriched with real analysis results by ReportService.
Produces a multi-slide deck covering: Executive Summary, KPIs, Data Quality,
Statistics, Correlations, Outliers, ML Models, Trends, and Recommendations.

Single Responsibility: converts ExecutiveReportData into a formatted PPTX deck.
"""

from __future__ import annotations

import os
from typing import Optional

from ai_engine.domain.entities import ExecutiveReportData

try:
    import pptx
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False


def _txt(paragraph, text: str, size=13, bold=False, color=None, italic=False):
    """Helper: add a run to a paragraph."""
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def _add_table(slide, headers, rows, left, top, width, height):
    """Add a styled table to a slide."""
    if not rows and not headers:
        return None
    col_count = len(headers)
    row_count = len(rows) + 1
    table_shape = slide.shapes.add_table(row_count, col_count, left, top, width, height)
    table = table_shape.table

    # Header row
    DARK = RGBColor(15, 23, 42)
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK

    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = str(val)[:80]
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9)
                p.font.color.rgb = RGBColor(30, 41, 59)
            if r_idx % 2 == 1:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(248, 250, 252)

    return table


class PPTXGenerator:
    """Generates a comprehensive PowerPoint presentation from ExecutiveReportData."""

    def __init__(self, data: ExecutiveReportData) -> None:
        self._data = data

    def generate(self, output_path: str) -> Optional[str]:
        if not PPTX_AVAILABLE:
            print("[PPTXGenerator] python-pptx is not installed.")
            return None

        try:
            prs = pptx.Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            blank = prs.slide_layouts[6]

            # Colors
            DARK = RGBColor(15, 23, 42)
            ACCENT = RGBColor(79, 70, 229)
            TEXT = RGBColor(30, 41, 59)
            MUTED = RGBColor(100, 116, 139)
            WHITE = RGBColor(255, 255, 255)
            CARD_BG = RGBColor(241, 245, 249)

            def banner(slide, title, subtitle):
                s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.4))
                s.fill.solid()
                s.fill.fore_color.rgb = DARK
                s.line.color.rgb = DARK
                tf = s.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                _txt(p, title, 24, True, WHITE)
                p2 = tf.add_paragraph()
                _txt(p2, subtitle, 13, False, RGBColor(203, 213, 225))

            def card(slide, left, top, w, h):
                c = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
                c.fill.solid()
                c.fill.fore_color.rgb = CARD_BG
                c.line.color.rgb = CARD_BG
                return c

            def textbox(slide, left, top, w, h):
                return slide.shapes.add_textbox(left, top, w, h).text_frame

            d = self._data

            # ======== SLIDE 1: COVER ========
            s1 = prs.slides.add_slide(blank)
            banner(s1, d.title, f"Dataset: {d.dataset_name} | {d.dataset_shape[0]:,} rows x {d.dataset_shape[1]} features")
            card(s1, Inches(1), Inches(2.0), Inches(11.333), Inches(4.8))
            tf = textbox(s1, Inches(1.3), Inches(2.2), Inches(10.7), Inches(4.4))
            tf.word_wrap = True
            p = tf.paragraphs[0]
            _txt(p, "Executive Analytics Briefing", 20, True, DARK)
            p.space_after = Pt(16)
            p2 = tf.add_paragraph()
            _txt(p2, d.summary_paragraph or f"Comprehensive analysis of {d.dataset_name}.", 14, False, TEXT)
            p2.line_spacing = 1.2

            # KPI cards on cover
            if d.kpis:
                p3 = tf.add_paragraph()
                p3.space_before = Pt(20)
                _txt(p3, "Key Metrics", 15, True, ACCENT)
                for kpi in d.kpis[:6]:
                    pk = tf.add_paragraph()
                    _txt(pk, f"  {kpi.get('title', '')}: {kpi.get('value', '0')}", 12, False, TEXT)

            # ======== SLIDE 2: DATA QUALITY ========
            s2 = prs.slides.add_slide(blank)
            banner(s2, "Data Quality Assessment", "Missing values, duplicates, and type consistency")
            card(s2, Inches(1), Inches(1.8), Inches(11.333), Inches(5.2))
            tf2 = textbox(s2, Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.8))
            tf2.word_wrap = True
            qd = d.quality_detail
            if qd:
                p = tf2.paragraphs[0]
                _txt(p, f"Quality Score: {qd.get('quality_score', d.quality_score):.1f}/100  |  Missing: {d.missing_percentage:.2f}%  |  Duplicates: {qd.get('duplicate_rows', 0):,}", 14, True, DARK)
                p.space_after = Pt(12)

                # Missing values table
                missing = qd.get("missing_per_column", [])
                if missing:
                    headers = ["Column", "Missing Count", "Missing %"]
                    rows = [[m["column"], str(m["missing_count"]), f"{m['missing_pct']:.2f}%"] for m in missing[:12]]
                    _add_table(s2, headers, rows, Inches(1.2), Inches(3.2), Inches(10.9), min(Inches(0.3) * (len(rows) + 1), Inches(3.5)))

                # Type issues
                type_issues = qd.get("type_issues", [])
                if type_issues:
                    y_offset = 3.2 + min((len(missing[:12]) + 1) * 0.3, 3.5) + 0.3
                    p_ti = tf2.add_paragraph()
                    _txt(p_ti, "Data Type Issues:", 12, True, ACCENT)
                    for ti in type_issues[:5]:
                        p_issue = tf2.add_paragraph()
                        _txt(p_issue, f"  - {ti['column']}: {ti['issue']}", 11, False, TEXT)
            else:
                p = tf2.paragraphs[0]
                _txt(p, f"Quality Score: {d.quality_score:.1f}/100  |  Missing: {d.missing_percentage:.2f}%", 14, True, DARK)

            # ======== SLIDE 3: STATISTICS ========
            s3 = prs.slides.add_slide(blank)
            banner(s3, "Descriptive Statistics", "Numeric and categorical column profiling")
            card(s3, Inches(1), Inches(1.8), Inches(5.5), Inches(5.2))
            card(s3, Inches(6.833), Inches(1.8), Inches(5.5), Inches(5.2))

            stats = d.statistics
            ns = stats.get("numeric_stats", [])
            if ns:
                headers = ["Column", "Mean", "Std", "Min", "Max", "Skew"]
                rows = [[s["column"], f"{s['mean']:.2f}", f"{s['std']:.2f}", f"{s['min']:.2f}", f"{s['max']:.2f}", f"{s['skewness']:.2f}"] for s in ns[:10]]
                _add_table(s3, headers, rows, Inches(1.1), Inches(2.0), Inches(5.3), min(Inches(0.3) * (len(rows) + 1), Inches(4.8)))
            else:
                tf3l = textbox(s3, Inches(1.3), Inches(2.2), Inches(5.0), Inches(1))
                _txt(tf3l.paragraphs[0], "No numeric columns found.", 12, False, MUTED)

            cs = stats.get("categorical_stats", [])
            if cs:
                headers = ["Column", "Unique", "Top Value", "Top Freq"]
                rows = []
                for c in cs[:10]:
                    top = c.get("top_values", [{}])[0] if c.get("top_values") else {}
                    rows.append([c["column"], str(c.get("unique", 0)), str(top.get("value", "N/A"))[:25], f"{top.get('frequency', 0):.1f}%"])
                _add_table(s3, headers, rows, Inches(6.933), Inches(2.0), Inches(5.3), min(Inches(0.3) * (len(rows) + 1), Inches(4.8)))
            else:
                tf3r = textbox(s3, Inches(7.1), Inches(2.2), Inches(5.0), Inches(1))
                _txt(tf3r.paragraphs[0], "No categorical columns found.", 12, False, MUTED)

            # ======== SLIDE 4: CORRELATIONS & OUTLIERS ========
            s4 = prs.slides.add_slide(blank)
            banner(s4, "Correlations & Outliers", "Feature relationships and anomaly detection")
            card(s4, Inches(1), Inches(1.8), Inches(5.5), Inches(5.2))
            card(s4, Inches(6.833), Inches(1.8), Inches(5.5), Inches(5.2))

            # Correlations
            pairs = d.correlations.get("top_pairs", []) if d.correlations else []
            if pairs:
                headers = ["Feature A", "Feature B", "Correlation"]
                rows = [[p["col_a"], p["col_b"], f"{p.get('correlation', 0):.4f}"] for p in pairs[:8]]
                _add_table(s4, headers, rows, Inches(1.1), Inches(2.0), Inches(5.3), min(Inches(0.3) * (len(rows) + 1), Inches(4.8)))
            else:
                tf4l = textbox(s4, Inches(1.3), Inches(2.2), Inches(5.0), Inches(1))
                _txt(tf4l.paragraphs[0], "No strong correlations found.", 12, False, MUTED)

            # Outliers
            iqr = d.outliers.get("iqr_outliers", []) if d.outliers else []
            if iqr:
                headers = ["Column", "Count", "%", "Lower", "Upper"]
                rows = [[o["column"], str(o.get("count", 0)), f"{o.get('percentage', 0):.1f}%", f"{o.get('lower_bound', 0):.2f}", f"{o.get('upper_bound', 0):.2f}"] for o in iqr[:8]]
                _add_table(s4, headers, rows, Inches(6.933), Inches(2.0), Inches(5.3), min(Inches(0.3) * (len(rows) + 1), Inches(4.8)))
            else:
                tf4r = textbox(s4, Inches(7.1), Inches(2.2), Inches(5.0), Inches(1))
                _txt(tf4r.paragraphs[0], "No outliers detected.", 12, False, MUTED)

            # ======== SLIDE 5: ML MODELS ========
            if d.ml_models:
                s5 = prs.slides.add_slide(blank)
                banner(s5, "ML Model Performance", f"{len(d.ml_models)} trained model(s)")
                card(s5, Inches(1), Inches(1.8), Inches(11.333), Inches(5.2))
                tf5 = textbox(s5, Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.8))
                tf5.word_wrap = True
                for m in d.ml_models:
                    pm = tf5.add_paragraph()
                    _txt(pm, f"{m['model_name']} ({m['algorithm']}) - {m.get('problem_type', '').title()}", 14, True, DARK)
                    pm.space_after = Pt(4)
                    metrics = m.get("metrics", {})
                    if metrics:
                        pm_metrics = tf5.add_paragraph()
                        metric_str = "  |  ".join([f"{k}: {v:.4f}" if isinstance(v, float) else f"{k}: {v}" for k, v in metrics.items()])
                        _txt(pm_metrics, metric_str, 11, False, MUTED)
                    pm_metrics.space_after = Pt(12)

            # ======== SLIDE 6: TRENDS & ANOMALIES ========
            if d.trends or (d.anomalies and d.anomalies.anomalous_rows_count > 0):
                s6 = prs.slides.add_slide(blank)
                banner(s6, "AI Trends & Anomalies", "Automated pattern detection")
                card(s6, Inches(1), Inches(1.8), Inches(5.5), Inches(5.2))
                card(s6, Inches(6.833), Inches(1.8), Inches(5.5), Inches(5.2))

                if d.trends:
                    headers = ["Column", "Direction", "R2", "Change"]
                    rows = [[t.column, t.direction.value.upper(), f"{t.r2_score:.3f}", f"{t.pct_change:+.1f}%"] for t in d.trends[:8]]
                    _add_table(s6, headers, rows, Inches(1.1), Inches(2.0), Inches(5.3), min(Inches(0.3) * (len(rows) + 1), Inches(4.8)))

                if d.anomalies and d.anomalies.anomalous_rows_count > 0:
                    headers = ["Row", "Column", "Value", "Reason"]
                    rows = [[str(item.row_index), item.column, str(item.value)[:20], item.reason[:40]] for item in d.anomalies.detected_anomalies[:8]]
                    _add_table(s6, headers, rows, Inches(6.933), Inches(2.0), Inches(5.3), min(Inches(0.3) * (len(rows) + 1), Inches(4.8)))

            # ======== SLIDE 7: RECOMMENDATIONS ========
            if d.recommendations:
                s7 = prs.slides.add_slide(blank)
                banner(s7, "Strategic Recommendations", "AI-generated action items")
                card(s7, Inches(1), Inches(1.8), Inches(11.333), Inches(5.2))
                tf7 = textbox(s7, Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.8))
                tf7.word_wrap = True
                for idx, r in enumerate(d.recommendations[:4], 1):
                    pr_ = tf7.add_paragraph()
                    impact_clr = RGBColor(239, 68, 68) if r.impact.value == "high" else (RGBColor(245, 158, 11) if r.impact.value == "medium" else RGBColor(34, 197, 94))
                    _txt(pr_, f"{idx}. {r.title}", 14, True, impact_clr)
                    pr_.space_after = Pt(2)
                    pcat = tf7.add_paragraph()
                    _txt(pcat, f"  {r.category} | {r.description[:120]}", 11, False, MUTED)
                    pcat.space_after = Pt(10)

            os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
            prs.save(output_path)
            return output_path
        except Exception as e:
            print(f"[PPTXGenerator] Error: {e}")
            import traceback
            traceback.print_exc()
            return None
